import os
import sys
import openpyxl
from tabulate import tabulate

# Configuración para soportar caracteres unicode de la tabla en Windows
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


def cargar_datos():
    """Carga los datos del archivo personas.xlsx en una lista."""
    # Buscar el archivo en la misma carpeta del script o en la carpeta actual
    ruta_archivo = os.path.join(os.path.dirname(__file__), "personas.xlsx")
    if not os.path.exists(ruta_archivo):
        ruta_archivo = "personas.xlsx"

    if not os.path.exists(ruta_archivo):
        print(f"Error: No se encontró el archivo '{ruta_archivo}'.")
        return []

    wb = openpyxl.load_workbook(ruta_archivo)
    hoja = wb.active

    datos = []
    # Recorremos todas las filas del archivo (a partir de la fila 1 de datos)
    for fila in range(1, hoja.max_row):
        fila_datos = [fila]  # Enumeración (#)
        for col in hoja.iter_cols(1, hoja.max_column):
            fila_datos.append(col[fila].value)
        datos.append(fila_datos)

    return datos


def mostrar_todos(datos):
    """Muestra todos los registros en formato de tabla profesional."""
    if not datos:
        print("No hay datos para mostrar.")
        return

    encabezados = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
    alineacion = ("center",) * 6
    print("\n--- TODOS LOS REGISTROS ---")
    print(tabulate(datos, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))


def reto1_nombres_emails(datos):
    """Reto 1: Mostrar únicamente Name y Email."""
    if not datos:
        print("No hay datos disponibles.")
        return

    # Extraer columnas: # (0), Name (2) y Email (4)
    filas_filtradas = []
    for fila in datos:
        filas_filtradas.append([fila[0], fila[2], fila[4]])

    encabezados = ["#", "Name", "Email"]
    alineacion = ("center", "left", "left")
    print("\n--- RETO 1: LISTADO DE NOMBRES Y CORREOS ---")
    print(tabulate(filas_filtradas, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))


def reto2_contar_registros(datos):
    """Reto 2: Mostrar cuántos registros tiene el archivo."""
    total = len(datos)
    print(f"\n--- RETO 2: CONTEO DE REGISTROS ---")
    print(f"El archivo contiene un total de {total} registros de personas.")


def reto3_filtrar_empresa(datos):
    """Reto 3: Filtrar personas cuya empresa contenga la palabra Group."""
    busqueda = input("\nIngrese término de búsqueda para la empresa (Enter para 'Group'): ").strip()
    if not busqueda:
        busqueda = "Group"

    coincidencias = []
    for fila in datos:
        empresa = str(fila[3]) if fila[3] is not None else ""
        if busqueda.lower() in empresa.lower():
            coincidencias.append(fila)

    print(f"\n--- RETO 3: EMPRESAS QUE CONTIENEN '{busqueda}' ---")
    if coincidencias:
        encabezados = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
        alineacion = ("center",) * 6
        print(tabulate(coincidencias, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))
        print(f"Total encontrados: {len(coincidencias)}")
    else:
        print(f"No se encontraron personas cuya empresa contenga '{busqueda}'.")


def reto4_buscar_por_id(datos):
    """Reto 4: Pedir un ID con input() y mostrar solo esa persona."""
    id_buscar = input("\nIngrese el ID de la persona a consultar: ").strip()

    if not id_buscar:
        print("Debe ingresar un ID válido.")
        return

    encontrado = None
    for fila in datos:
        id_actual = str(fila[1]).strip()
        if id_actual.lower() == id_buscar.lower():
            encontrado = fila
            break

    print(f"\n--- RETO 4: RESULTADO DE BÚSQUEDA POR ID ---")
    if encontrado:
        encabezados = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
        alineacion = ("center",) * 6
        print(tabulate([encontrado], headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))
    else:
        print(f"No se encontró ningún registro con el ID: '{id_buscar}'.")


def menu():
    """Reto 5: Menú interactivo con opciones organizadas."""
    datos = cargar_datos()

    while True:
        print("\n" + "=" * 42)
        print("      MENÚ - TALLER PRÁCTICO PYTHON")
        print("=" * 42)
        print("1. Ver todos los registros")
        print("2. Buscar persona por ID (Reto 4)")
        print("3. Filtrar por empresa 'Group' (Reto 3)")
        print("4. Contar cantidad de registros (Reto 2)")
        print("5. Mostrar solo Nombre y Email (Reto 1)")
        print("6. Salir")
        print("=" * 42)

        opcion = input("Seleccione una opción (1-6): ").strip()

        if opcion == "1":
            mostrar_todos(datos)
        elif opcion == "2":
            reto4_buscar_por_id(datos)
        elif opcion == "3":
            reto3_filtrar_empresa(datos)
        elif opcion == "4":
            reto2_contar_registros(datos)
        elif opcion == "5":
            reto1_nombres_emails(datos)
        elif opcion == "6":
            print("\nSaliendo del programa. ¡Hasta pronto!")
            break
        else:
            print("\nOpción inválida. Ingrese un número entre 1 y 6.")


if __name__ == "__main__":
    menu()
